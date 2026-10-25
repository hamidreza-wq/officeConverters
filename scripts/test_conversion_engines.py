"""Real fixtures and conversion output assertions; no network services or mocked engines."""

from __future__ import annotations

import csv
import io
import json
import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from convert_document import CAPABILITIES, Engine, RejectedDocument
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def settings() -> dict:
    return {
        "libreoffice": os.environ.get("CONVERSION_LIBREOFFICE", r"C:\Program Files\LibreOffice\program\soffice.com"),
        "tesseract": os.environ.get("CONVERSION_TESSERACT", str(Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Tesseract-OCR" / "tesseract.exe")),
        "timeout": 120,
    }


def partial_grid_pdf(path: Path, *, header_only: bool = False) -> None:
    writer = PdfWriter()
    for page_number in range(2):
        page = writer.add_blank_page(width=595, height=842)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        commands = []
        boundaries = [57, 116, 264, 303, 382, 460, 538]
        headers = ["Date", "Description", "Type", "Money In", "Money Out", "Balance"]
        for left, right in zip(boundaries, boundaries[1:]):
            commands.append(f"{left} 650 m {right} 650 l S {left} 625 m {right} 625 l S")
        for x in boundaries[1:-1]:
            commands.append(f"{x} 650 m {x} 625 l S")
        for x, label in zip(boundaries, headers):
            commands.append(f"BT /F1 0.05 Tf 1 g {x + 6} 640 Td (Column) Tj ET")
            commands.append(f"BT /F1 9 Tf 0 g {x + 6} 634 Td ({label}) Tj ET")
        if not header_only:
            rows = [
                [f"0{page_number + 1} Sep", "Sample payment", "DEB", "", "12.50", "987.50"],
                [f"0{page_number + 1} Sep", "Sample credit", "CR", "100.00", "", "1087.50"],
            ]
            for index, row in enumerate(rows):
                top = 625 - index * 30
                for x in boundaries[1:-1]:
                    commands.append(f"{x} {top} m {x} {top - 30} l S")
                for x, label in zip(boundaries, row):
                    commands.append(f"BT /F1 0.05 Tf 1 g {x + 6} {top - 8} Td (hidden label) Tj ET")
                    commands.append(f"BT /F1 9 Tf 0 g {x + 6} {top - 16} Td ({label}) Tj ET")
                if index == 0:
                    commands.append(f"BT /F1 9 Tf 0 g 122 {top - 26} Td (continued description) Tj ET")
        commands.append("BT /F1 9 Tf 57 530 Td (Footer must not become a transaction) Tj ET")
        stream = DecodedStreamObject()
        stream.set_data("\n".join(commands).encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    writer.write(path)
    writer.close()


def fixtures(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    partial_grid_pdf(folder / "partial-grid.pdf")
    document = Document()
    document.add_heading("Engine test", level=1)
    document.add_paragraph("Sample document 123.4500")
    document.save(folder / "sample.docx")
    (folder / "sample.txt").write_text("Engine test\nSample document 123.4500", encoding="utf-8")
    (folder / "sample.rtf").write_text(r"{\rtf1\ansi Engine test\par Sample document 123.4500}", encoding="ascii")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Engine test"
    slide.placeholders[1].text = "Sample document 123.4500"
    presentation.save(folder / "sample.pptx")
    book = Workbook()
    sheet = book.active
    sheet.title = "First"
    sheet.append(["Name", "Amount", "Reference"])
    sheet.append(["Sample", "123.4500", "000123"])
    sheet.append(["Other", "42.00", "99999999999999999999"])
    book.save(folder / "sample.xlsx")
    book.close()
    complex_book = Workbook()
    first = complex_book.active
    first.title = "Revenue"
    first.append(["Item", "Amount"])
    first.append(["Sample", 12])
    first.append(["Other", 3])
    first.append(["Total", "=SUM(B2:B3)"])
    first.merge_cells("D1:E1")
    first["D1"] = "Merged header"
    second = complex_book.create_sheet("Expenses")
    second.append(["Item", "Amount"])
    second.append(["Cost", 2])
    complex_book.save(folder / "complex.xlsx")
    complex_book.close()

    image = Image.new("RGB", (1200, 500), "white")
    draw = ImageDraw.Draw(image)
    font_path = r"C:\Windows\Fonts\arial.ttf" if os.name == "nt" else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    font = ImageFont.truetype(font_path, 44)
    for y, left, right in [(50, "Name", "Amount"), (170, "Sample", "123"), (290, "Other", "456")]:
        draw.text((60, y), left, font=font, fill="black")
        draw.text((650, y), right, font=font, fill="black")
    for ext, fmt in [("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"), ("webp", "WEBP"), ("bmp", "BMP"), ("tiff", "TIFF"), ("gif", "GIF")]:
        image.save(folder / f"sample.{ext}", fmt)
    image.save(folder / "scanned.pdf", "PDF", resolution=150)
    image.save(folder / "animated.gif", "GIF", save_all=True, append_images=[Image.new("RGB", image.size, "lightblue")], duration=100)
    image.close()

    writer = PdfWriter()
    page = writer.add_blank_page(width=360, height=240)
    font_object = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font_object)})})
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 30 210 Td (Engine test) Tj ET\n"
        b"0.5 w 30 60 300 120 re S 30 100 m 330 100 l S 30 140 m 330 140 l S 190 60 m 190 180 l S\n"
        b"BT /F1 12 Tf 40 155 Td (Name) Tj 160 0 Td (Amount) Tj ET\n"
        b"BT /F1 12 Tf 40 115 Td (Sample) Tj 160 0 Td (123) Tj ET\n"
        b"BT /F1 12 Tf 40 75 Td (Other) Tj 160 0 Td (456) Tj ET\n"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    with (folder / "sample.pdf").open("wb") as output:
        writer.write(output)
    writer.close()
    engine = Engine(settings())
    for source, target in [("docx", "odt"), ("pptx", "odp"), ("xlsx", "ods")]:
        converted = engine.office(folder / f"sample.{source}", source, target, folder / f"prepare-{target}")
        (folder / f"sample.{target}").write_bytes(converted.read_bytes())


def verify_output(path: Path) -> None:
    assert path.is_file() and path.stat().st_size > 0
    extension = path.suffix.lower()
    if extension == ".pdf":
        reader = PdfReader(path)
        assert len(reader.pages) > 0
    elif extension == ".docx":
        doc = Document(path)
        assert "Sample" in "\n".join(paragraph.text for paragraph in doc.paragraphs)
    elif extension == ".pptx":
        assert len(Presentation(path).slides) > 0
    elif extension == ".xlsx":
        book = load_workbook(path, data_only=False)
        assert book.sheetnames
        assert any(cell.value is not None for row in book.active for cell in row)
        book.close()
    elif extension in {".odt", ".odp", ".ods"}:
        with zipfile.ZipFile(path) as package:
            assert "content.xml" in package.namelist()
    elif extension in {".png", ".jpg", ".webp", ".bmp", ".tiff"}:
        with Image.open(path) as image:
            image.verify()
    elif extension == ".zip":
        with zipfile.ZipFile(path) as package:
            assert package.namelist()
            assert package.testzip() is None
    elif extension == ".rtf":
        assert path.read_bytes().startswith(b"{\\rtf")
    elif extension in {".txt", ".csv", ".tsv", ".json"}:
        text = path.read_text(encoding="utf-8-sig")
        assert "Sample" in text


class ConversionEnginesTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.temp = tempfile.TemporaryDirectory(prefix="engine-test-")
        cls.folder = Path(cls.temp.name)
        fixtures(cls.folder)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.temp.cleanup()

    def test_every_advertised_local_pair_produces_a_readable_output(self) -> None:
        count = 0
        for group in CAPABILITIES.values():
            for source, targets in group.items():
                for target in targets:
                    with self.subTest(source=source, target=target):
                        directory = self.folder / f"test-{source}-{target}"
                        result, preview = Engine(settings()).convert(self.folder / f"sample.{source}", source, target, directory)
                        verify_output(result)
                        self.assertLessEqual(len(preview), 10)
                        count += 1
        self.assertEqual(count, sum(len(targets) for group in CAPABILITIES.values() for targets in group.values()))

    def test_complex_spreadsheet_preserves_every_sheet_and_calculates_formulas(self) -> None:
        directory = self.folder / "complex-export"
        engine = Engine(settings())
        result, preview = engine.convert(self.folder / "complex.xlsx", "xlsx", "csv", directory)
        self.assertEqual(".zip", result.suffix)
        with zipfile.ZipFile(result) as package:
            self.assertEqual(2, len(package.namelist()))
            values = [package.read(name).decode() for name in package.namelist()]
        self.assertTrue(any("Total,15" in text for text in values))
        self.assertTrue(any("Cost,2" in text for text in values))
        self.assertTrue(any("worksheet" in warning for warning in engine.messages))

    def test_csv_preview_matches_neutralized_output_values(self) -> None:
        source = self.folder / "formula-like.json"
        source.write_text(json.dumps([["Name", "Value"], ["Sample", "=1+1"]]), encoding="utf-8")
        output, preview = Engine(settings()).convert(source, "json", "csv", self.folder / "formula-preview")
        self.assertEqual("'=1+1", preview[1][1])
        self.assertIn("'=1+1", output.read_text(encoding="utf-8"))

    def test_partial_grid_pdf_exports_data_rows_and_preserves_empty_columns(self) -> None:
        source = self.folder / "partial-grid.pdf"
        partial_grid_pdf(source)
        for target in ["csv", "xlsx"]:
            with self.subTest(target=target):
                output, preview = Engine(settings()).convert(source, "pdf", target, self.folder / f"partial-grid-{target}")
                if target == "csv":
                    with zipfile.ZipFile(output) as archive:
                        tables = [list(csv.reader(io.StringIO(archive.read(name).decode("utf-8")))) for name in archive.namelist()]
                else:
                    book = load_workbook(output)
                    tables = [[[value or "" for value in row] for row in sheet.values] for sheet in book]
                    book.close()
                self.assertEqual(2, len(tables))
                for rows in tables:
                    self.assertEqual(3, len(rows))
                    self.assertEqual(["Date", "Description", "Type", "Money In", "Money Out", "Balance"], rows[0])
                    self.assertEqual(["", "12.50", "987.50"], rows[1][3:])
                    self.assertEqual(["100.00", "", "1087.50"], rows[2][3:])
                    self.assertIn("continued description", rows[1][1])
                    self.assertNotIn("hidden", str(rows))
                    self.assertNotIn("Footer", str(rows))
                self.assertEqual(tables[0], preview)

    def test_header_only_pdf_is_not_reported_as_a_successful_table_conversion(self) -> None:
        source = self.folder / "header-only.pdf"
        partial_grid_pdf(source, header_only=True)
        with self.assertRaisesRegex(RejectedDocument, "No reliable table"):
            Engine(settings()).convert(source, "pdf", "csv", self.folder / "header-only-result")

    def test_fully_ruled_pdf_table_still_exports_exact_values(self) -> None:
        expected = [["Name", "Amount"], ["Sample", "123"], ["Other", "456"]]
        for target in ["csv", "xlsx"]:
            with self.subTest(target=target):
                output, preview = Engine(settings()).convert(self.folder / "sample.pdf", "pdf", target, self.folder / f"ruled-{target}")
                self.assertEqual(expected, preview)
                if target == "csv":
                    with output.open(newline="", encoding="utf-8") as stream:
                        self.assertEqual(expected, list(csv.reader(stream)))
                else:
                    book = load_workbook(output)
                    self.assertEqual(expected, [list(row) for row in book.active.values])
                    book.close()

    def test_scanned_pdf_table_ocr_and_searchable_pdf(self) -> None:
        engine = Engine(settings())
        output, preview = engine.convert(self.folder / "scanned.pdf", "pdf", "xlsx", self.folder / "scanned-table")
        self.assertTrue(any("Sample" in cell for row in preview for cell in row))
        self.assertTrue(any("123" in cell for row in preview for cell in row))
        verify_output(output)
        output, _ = Engine(settings()).convert(self.folder / "scanned.pdf", "pdf", "ocrpdf", self.folder / "ocr-pdf")
        self.assertIn("Sample", PdfReader(output).pages[0].extract_text())

    def test_every_animation_frame_is_retained(self) -> None:
        output, _ = Engine(settings()).convert(self.folder / "animated.gif", "gif", "png", self.folder / "animated")
        with zipfile.ZipFile(output) as package:
            self.assertEqual(2, len(package.namelist()))

    def test_archives_convert_all_entries_and_reject_unsafe_ones(self) -> None:
        path = self.folder / "documents.zip"
        with zipfile.ZipFile(path, "w") as package:
            package.write(self.folder / "sample.docx", "folder/document.docx")
            package.write(self.folder / "sample.png", "image.png")
        output, preview = Engine(settings()).convert(path, "zip", "pdf", self.folder / "archive-results")
        with zipfile.ZipFile(output) as package:
            self.assertEqual(3, len(package.namelist()))
            for name in package.namelist():
                if name.endswith(".pdf"):
                    self.assertGreater(len(PdfReader(io.BytesIO(package.read(name))).pages), 0)
        self.assertEqual(3, len(preview))
        for name in ["../escape.docx", "C:/escape.docx", "nested.zip", "script.exe"]:
            with self.subTest(name=name):
                with zipfile.ZipFile(path, "w") as package:
                    package.writestr(name, b"unsafe")
                with self.assertRaises(RejectedDocument):
                    Engine(settings()).convert(path, "zip", "pdf", self.folder / "unsafe-archive")

    def test_every_advertised_archive_target_has_a_real_compatible_conversion(self) -> None:
        sources = {
            "csv": "xlsx", "tsv": "xlsx", "json": "xlsx", "xlsx": "ods", "ods": "xlsx",
            "pdf": "docx", "docx": "odt", "odt": "docx", "rtf": "docx", "txt": "docx",
            "pptx": "odp", "odp": "pptx", "png": "jpg", "jpg": "png", "webp": "png",
            "bmp": "png", "tiff": "png", "ocrpdf": "png",
        }
        advertised = set(["csv", "tsv", "json", "xlsx"])
        for group in CAPABILITIES.values():
            for targets in group.values():
                advertised.update(targets)
        self.assertEqual(advertised, set(sources))
        for target, source in sources.items():
            with self.subTest(target=target):
                archive_path = self.folder / f"batch-{target}.zip"
                with zipfile.ZipFile(archive_path, "w") as archive:
                    archive.write(self.folder / f"sample.{source}", f"document.{source}")
                output, _ = Engine(settings()).convert(archive_path, "zip", target, self.folder / f"batch-result-{target}")
                with zipfile.ZipFile(output) as archive:
                    files = [name for name in archive.namelist() if not name.endswith("manifest.json")]
                    self.assertEqual(1, len(files))
                    local = self.folder / f"inspect-{target}" / Path(files[0]).name
                    local.parent.mkdir(exist_ok=True)
                    local.write_bytes(archive.read(files[0]))
                    verify_output(local)

    def test_encrypted_pdf_external_links_expansion_and_pixel_limits(self) -> None:
        encrypted = self.folder / "encrypted.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.encrypt("password")
        writer.write(encrypted)
        writer.close()
        with self.assertRaises(RejectedDocument):
            Engine(settings()).convert(encrypted, "pdf", "txt", self.folder / "encrypted-result")
        with self.assertRaises(RejectedDocument):
            engine = Engine({**settings(), "max_uncompressed_bytes": 100})
            engine.office_preflight(self.folder / "sample.docx", "docx")
        external = self.folder / "external.docx"
        with zipfile.ZipFile(external, "w") as package:
            package.writestr("_rels/.rels", '<Relationships><Relationship TargetMode="External" Target="https://example.test"/></Relationships>')
        with self.assertRaises(RejectedDocument):
            Engine(settings()).office_preflight(external, "docx")


if __name__ == "__main__":
    if len(sys.argv) == 3 and sys.argv[1] == "--fixtures":
        fixtures(Path(sys.argv[2]))
    else:
        unittest.main()
